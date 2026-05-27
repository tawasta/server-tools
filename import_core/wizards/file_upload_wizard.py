import base64
import csv
import io
import os

from odoo import _, exceptions, fields, models

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


class FileUploadWizard(models.TransientModel):
    _name = "file.upload.wizard"
    _description = "File Upload Wizard"

    template_id = fields.Many2one(
        comodel_name="generic.import.template",
        string="Template",
    )
    file_data = fields.Binary("File")
    file_name = fields.Char("Filename")

    def _read_rows_from_csv(self, file_bytes: bytes) -> list[dict]:
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1")

        file_stream = io.StringIO(text)
        return list(csv.DictReader(file_stream, delimiter=","))

    def _read_rows_from_xlsx(self, file_bytes: bytes) -> list[dict]:
        if load_workbook is None:
            raise exceptions.UserError(
                _("XLSX import requires python dependency 'openpyxl'.")
            )

        wb = load_workbook(
            filename=io.BytesIO(file_bytes),
            data_only=True,
            read_only=True,
        )
        ws = wb.worksheets[0]

        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return []

        headers = []
        for h in header_row:
            headers.append(str(h).strip() if h is not None else "")

        if not any(headers):
            raise exceptions.UserError(
                _("XLSX file has no header row (first row is empty).")
            )

        result = []
        for row in rows_iter:
            if row is None:
                continue

            data = {}
            any_value = False

            for idx, col_name in enumerate(headers):
                if not col_name:
                    continue

                val = row[idx] if idx < len(row) else None
                if val is None:
                    data[col_name] = ""
                    continue

                if isinstance(val, str):
                    data[col_name] = val.strip()
                else:
                    data[col_name] = str(val).strip()

                if data[col_name] not in ("", None, False):
                    any_value = True

            if any_value:
                result.append(data)

        return result

    def _is_xlsx(self, file_bytes: bytes) -> bool:
        return file_bytes[:4] == b"PK\x03\x04"

    def _read_rows_from_file(self) -> list[dict]:
        self.ensure_one()

        if not self.file_data:
            raise exceptions.UserError(_("Tiedostoa ei ole ladattu."))

        file_bytes = base64.b64decode(self.file_data)

        if self._is_xlsx(file_bytes):
            return self._read_rows_from_xlsx(file_bytes)

        name = (self.file_name or "").strip().lower()
        root, ext = os.path.splitext(name)
        if ext == ".xlsx":
            return self._read_rows_from_xlsx(file_bytes)

        return self._read_rows_from_csv(file_bytes)

    def _split_lines_by_model(self, row: dict) -> dict:
        self.ensure_one()
        result = {}

        for col_name, cell_value in row.items():
            tlines = self.template_id.line_ids.filtered(
                lambda li, col=col_name: li.csv_column_name == col
            )

            for line in tlines:
                model_name = line.field_id.model
                field_name = line.field_id.name

                if model_name not in self.env.registry:
                    continue

                result.setdefault(model_name, {"create": [], "search": []})
                result[model_name]["create"].append((field_name, cell_value))

                if line.is_search_field and cell_value not in ("", None, False):
                    result[model_name]["search"].append((field_name, cell_value))

        result["_state_links"] = [
            {
                "apply_on": link.apply_on,
                "target_model": link.target_model_id.model,
                "target_field": link.target_field_id.name,
                "target_state_key": link.target_state_key,
                "source_state_key": link.source_state_key,
            }
            for link in self.template_id.state_link_ids
        ]

        return result

    def _apply_state_links(self, state):
        self.ensure_one()

        for link in self.template_id.state_link_ids.sorted(
            key=lambda l: (l.sequence, l.id)
        ):
            if link.apply_on != "write":
                continue

            target_record = state.get(link.target_state_key)
            source_record = state.get(link.source_state_key)

            if not target_record or not source_record:
                continue

            if link.target_field_id.ttype != "many2one":
                continue

            if target_record._name != link.target_model_id.model:
                continue

            field_name = link.target_field_id.name

            current_value = target_record[field_name]
            if current_value and current_value.id == source_record.id:
                continue

            target_record.write({
                field_name: source_record.id,
            })

    def create_records_from_file(self):
        self.ensure_one()

        if not self.template_id:
            raise exceptions.UserError(_("Template is required."))

        rows = self._read_rows_from_file()
        if not rows:
            raise exceptions.UserError(_("The uploaded file contains no data rows."))

        step_lines = self.template_id.step_line_ids.sorted(
            key=lambda line: (line.sequence, line.id)
        )
        if not step_lines:
            raise exceptions.UserError(
                _("Templatelta puuttuu Import steps -määrittely.")
            )

        state = {}

        for row_index, row in enumerate(rows, start=2):
            if not any(row.values()):
                continue

            state.pop("_skip_rest", None)
            lines_by_model = self._split_lines_by_model(row)

            for step_line in step_lines:
                step = step_line.step_id

                if not step.models_installed():
                    continue

                state = (
                    self.env["generic.import.runner"].run_step(
                        step.code,
                        row_index,
                        row,
                        lines_by_model,
                        state,
                    )
                    or state
                )

                self._apply_state_links(state)

                if state.get("_skip_rest"):
                    break

        return {
            "type": "ir.actions.client",
            "tag": "reload",
        }